import io
import subprocess
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

from django.test import SimpleTestCase

from . import xlsx
from .xlsx import Cell, Sheet

try:
    import openpyxl
except ImportError:                      # the writer needs nothing; only these checks use openpyxl to read the file back
    openpyxl = None


def load(sheets):
    return openpyxl.load_workbook(io.BytesIO(xlsx.build_workbook(sheets)))


@unittest.skipIf(openpyxl is None, "openpyxl is only needed to read the file back in these checks")
class WriterTests(SimpleTestCase):
    def test_values_formats_and_styles_survive_a_round_trip(self):
        rows = [[Cell("Cost report", xlsx.TITLE)], [],
                [Cell("Site", xlsx.HEADER), Cell("Days", xlsx.HEADER), Cell("Wages", xlsx.HEADER), Cell("From", xlsx.HEADER)],
                ["P1 Tower A", Cell(6, xlsx.INT), Cell(Decimal("30.000"), xlsx.MONEY), Cell(date(2026, 3, 1), xlsx.DATE)],
                [Cell("Total", xlsx.TEXT_BOLD), Cell(6, xlsx.INT_BOLD), Cell(Decimal("1234.5"), xlsx.MONEY_BOLD), None]]
        ws = load([Sheet("Cost", rows, widths=[24, 10, 14, 14], freeze_row=3)])["Cost"]
        self.assertEqual(ws["A1"].value, "Cost report")
        self.assertTrue(ws["A1"].font.b and ws["A1"].font.sz == 14)
        self.assertEqual([c.value for c in ws[3]], ["Site", "Days", "Wages", "From"])
        self.assertTrue(ws["A3"].font.b)
        self.assertEqual((ws["B4"].value, ws["B4"].number_format), (6, "#,##0"))
        self.assertEqual((ws["C4"].value, ws["C4"].number_format), (30, "#,##0.000"))
        self.assertEqual(ws["D4"].value.date(), date(2026, 3, 1))
        self.assertIn("mmm", ws["D4"].number_format)
        self.assertEqual((ws["C5"].value, ws["C5"].font.b), (1234.5, True))
        self.assertEqual(ws.column_dimensions["A"].width, 24)
        self.assertEqual(ws.freeze_panes, "A4")

    def test_text_is_never_run_as_a_formula_and_odd_characters_are_safe(self):
        nasty = ["=1+1", "+SUM(A1)", "@x", '<b>&"quotes"</b>', "عبدالله محمد", "line\nbreak", "bell\x07gone", "  padded  "]
        ws = load([Sheet("T", [[v] for v in nasty])])["T"]
        got = [ws.cell(row=i, column=1).value for i in range(1, len(nasty) + 1)]
        self.assertEqual(got[:6], nasty[:6])
        self.assertEqual(got[6], "bellgone")
        self.assertEqual(got[7], "  padded  ")
        self.assertEqual(ws["A1"].data_type, "s")                       # stored as text, not "f" for formula

    def test_sheet_names_are_made_valid_and_unique(self):
        sheets = [Sheet("By site/trade: A*B?", [["x"]]), Sheet("By site/trade: A*B?", [["y"]]), Sheet("", [["z"]]),
                  Sheet("A very long sheet name that goes past thirty one characters", [["w"]])]
        names = load(sheets).sheetnames
        self.assertEqual(len(set(n.lower() for n in names)), 4)
        for n in names:
            self.assertLessEqual(len(n), 31)
            self.assertFalse(set(n) & set("[]:*?/\\"))

    def test_empty_cells_and_wide_sheets(self):
        row = [Cell(i, xlsx.INT) for i in range(30)]
        ws = load([Sheet("Wide", [row, [None] * 30])])["Wide"]
        self.assertEqual(ws["AD1"].value, 29)                          # column 30 is AD
        self.assertEqual(xlsx._column(0), "A")
        self.assertEqual((xlsx._column(25), xlsx._column(26), xlsx._column(701), xlsx._column(702)), ("Z", "AA", "ZZ", "AAA"))

    def test_booleans_and_none(self):
        ws = load([Sheet("B", [[True, False, None, "", 0]])])["B"]
        self.assertEqual([ws.cell(row=1, column=c).value for c in (1, 2, 3, 4, 5)], ["Yes", "No", None, None, 0])


@unittest.skipIf(openpyxl is None or not Path("/usr/bin/soffice").exists(), "needs openpyxl and LibreOffice to open the file like Excel would")
class OpensInARealSpreadsheetProgramTests(SimpleTestCase):
    def test_libreoffice_opens_it_without_repair_and_reads_the_numbers(self):
        data = xlsx.build_workbook([Sheet("Report", [[Cell("Title", xlsx.TITLE)], [Cell("A", xlsx.HEADER), Cell("B", xlsx.HEADER)],
                                                      ["x", Cell(Decimal("12.345"), xlsx.MONEY)], ["y", Cell(Decimal("7.5"), xlsx.MONEY)]],
                                           widths=[20, 12], freeze_row=2)])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "r.xlsx"
            path.write_bytes(data)
            done = subprocess.run(["soffice", "--headless", "--convert-to", "csv", "--outdir", tmp, str(path)],
                                  capture_output=True, text=True, timeout=120)
            csv = (Path(tmp) / "r.csv")
            self.assertTrue(csv.exists(), done.stderr)
            self.assertEqual(csv.read_text().splitlines(), ["Title,", "A,B", "x,12.345", "y,7.5"])
