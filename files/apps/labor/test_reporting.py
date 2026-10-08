import io
import unittest
from datetime import date
from decimal import Decimal

from django.contrib.auth import get_user_model

from apps.attendance.models import Attendance
from apps.employees.models import Employee
from apps.scheduling.models import Holiday

from . import deployment, reporting, services, sitesheet, timekeeping
from . import overtime_services as ot
from .allocation import WorkOrder
from .overtime import OvertimeClaim
from .test_overtime import OvertimeCase
from .test_sitesheet import FRI, SAT, SUN
from .test_timekeeping import MON, TUE

D = Decimal
WED, THU = date(2026, 3, 4), date(2026, 3, 5)

try:
    import openpyxl
except ImportError:
    openpyxl = None


class ReportCase(OvertimeCase):
    """Ravi (direct, 5.000 a day) has worked 1-4, 5 and 7 March at P1 / Site One. Everything else is built per test."""

    def setUp(self):
        super().setUp()
        self.user = self.root

    def contracted(self, name="Imran", contractor="default", rate="6", project=None, location=None, start=date(2026, 1, 5)):
        contractor = self.contractor if contractor == "default" else contractor
        e = Employee(company=self.co, first_name=name, joining_date=date(2026, 1, 1))
        emp = services.create_labor_worker(e, engagement="contracted", contractor=contractor, trade=self.mason,
                                           wage_basis="daily", rate=D(rate)).employee
        deployment.allocate(emp, project=project or self.p1, location=location or self.site1, effective_from=start)
        return emp

    def present(self, *employees, project=None, location=None, days=(SUN, SAT)):
        sitesheet.save_sheet(self.root, project or self.p1, location or self.site1, days[0], days[1], {}, "present")

    def approve_ravi_overtime(self):
        self.confirmed({self.hkey(self.emp, MON): "10", self.hkey(self.emp, TUE): "11"})
        self.prepare()
        ot.decide(self.root, OvertimeClaim.objects.filter(date=MON), True)         # Monday approved, Tuesday pending

    def cost(self, group="site", first=SUN, last=SAT, **kw):
        return reporting.cost_report(self.user, first, last, group, **kw)

    def only(self, report, index=0):
        return report.tables[index]

    def row(self, table, label):
        return next(r for r in table.rows if label in r)


class DeploymentTests(ReportCase):
    def test_counts_by_site_split_into_direct_and_contracted(self):
        self.contracted("Imran")
        self.contracted("Salim", contractor=None)
        deployment.allocate(self.make_other("Far"), project=self.p2, location=self.site2, effective_from=date(2026, 1, 5))
        t = self.only(reporting.deployment_report(self.user, date(2026, 3, 3)))
        self.assertEqual(self.row(t, "Site One")[2:], [1, 2, 3])
        self.assertEqual(self.row(t, "Site Two")[2:], [1, 0, 1])
        self.assertEqual(t.total[2:], [2, 2, 4])

    def test_trade_and_contractor_tables_against_sites(self):
        self.contracted("Imran")
        self.contracted("Salim", contractor=None)
        rep = reporting.deployment_report(self.user, date(2026, 3, 3))
        trades, contractors = rep.tables[1], rep.tables[2]
        self.assertEqual([c.label for c in trades.columns], ["Trade", "P1 / Site One", "Total"])
        self.assertEqual(self.row(trades, "Mason"), ["Mason", 3, 3])
        names = [r[0] for r in contractors.rows]
        self.assertEqual(names, ["Gulf Manpower", "Contracted, engaged directly", "Direct employees"])
        self.assertEqual(contractors.total, ["Total", 3, 3])

    def test_the_date_decides_who_was_where(self):
        self.assertEqual(self.only(reporting.deployment_report(self.user, date(2026, 1, 4))).rows, [])      # before Ravi was placed
        deployment.allocate(self.emp, project=self.p2, location=self.site2, effective_from=date(2026, 3, 10))
        before = self.only(reporting.deployment_report(self.user, date(2026, 3, 9)))
        after = self.only(reporting.deployment_report(self.user, date(2026, 3, 10)))
        self.assertEqual([r[1] for r in before.rows], ["Site One"])
        self.assertEqual([r[1] for r in after.rows], ["Site Two"])

    def test_people_who_left_are_not_counted_and_the_note_says_so(self):
        Employee.objects.filter(pk=self.emp.pk).update(status="separated")
        rep = reporting.deployment_report(self.user, date(2026, 3, 3))
        self.assertEqual(self.only(rep).rows, [])
        self.assertTrue(any("left the company" in n for n in rep.notes))

    def test_workers_not_on_any_site_are_noted_for_today_only(self):
        self.make_other("Idle")
        today = reporting.deployment_report(self.user, date.today())
        self.assertTrue(any("1 worker is not on any site" in n for n in today.notes))
        past = reporting.deployment_report(self.user, date(2026, 3, 3))
        self.assertFalse(any("not on any site" in n for n in past.notes))

    def test_other_companies_see_nothing(self):
        outsider = get_user_model().objects.create_user("out")
        self.assertEqual(self.only(reporting.deployment_report(outsider, date(2026, 3, 3))).rows, [])


class ManpowerTests(ReportCase):
    def test_people_per_day_and_man_days(self):
        other = self.contracted("Imran")
        self.present()                                                        # both present Sun to Thu and Sat (Friday is off)
        Attendance.objects.filter(employee=other, date=MON).update(status="half_day")
        Attendance.objects.filter(employee=other, date=TUE).update(status="absent")
        t = self.only(reporting.manpower_report(self.user, SUN, SAT))
        row = t.rows[0]
        self.assertEqual(row[0], "P1 / Site One")
        self.assertEqual(row[1:8], [2, 2, 1, 2, 2, 0, 2])                      # Sun Mon Tue Wed Thu Fri Sat; a half day is one person
        self.assertEqual(row[8], D("10.5"))                                    # 6 + 5 present days... half day counted as half
        self.assertEqual(t.total[1:], [2, 2, 1, 2, 2, 0, 2, D("10.5")])
        self.assertEqual([c.label[:2] for c in t.columns[1:3]], ["1 ", "2 "])

    def test_records_without_a_site_are_listed_apart_and_staff_are_left_out(self):
        Attendance.objects.filter(employee=self.emp, date=SUN).update(project=None, location=None)
        staff = Employee.objects.create(company=self.co, first_name="Office", employee_no="S-1", joining_date=date(2026, 1, 1))
        Attendance.objects.create(employee=staff, date=SUN, status="present", project=self.p1, location=self.site1)
        names = [r[0] for r in self.only(reporting.manpower_report(self.user, SUN, SAT)).rows]
        self.assertEqual(names, ["P1 / Site One", "(no site recorded)"])
        self.assertEqual(self.only(reporting.manpower_report(self.user, SUN, SUN)).rows[0][0], "(no site recorded)")


class CostTests(ReportCase):
    def test_a_daily_workers_wages_hours_and_overtime(self):
        self.approve_ravi_overtime()
        t = self.only(self.cost())
        row = t.rows[0]
        self.assertEqual(row[:3], ["P1", "Tower A", "Site One"])
        workers, days, hours, wages, ot_hours, ot_ok, total, ot_pending = row[3:]
        self.assertEqual((workers, days, hours), (1, D("6"), D("53")))
        self.assertEqual((wages, ot_hours, ot_ok, total, ot_pending), (D("30.000"), D("2"), D("1.563"), D("31.563"), D("2.344")))
        self.assertEqual(t.total[3:], [1, D("6"), D("53"), D("30.000"), D("2"), D("1.563"), D("31.563"), D("2.344")])

    def test_half_days_pay_half_and_absence_and_leave_pay_nothing(self):
        Attendance.objects.filter(employee=self.emp, date=SUN).update(status="half_day")
        Attendance.objects.filter(employee=self.emp, date=MON).update(status="absent")
        Attendance.objects.filter(employee=self.emp, date=TUE).update(status="paid_leave")
        row = self.only(self.cost()).rows[0]
        self.assertEqual((row[4], row[6]), (D("3.5"), D("17.500")))             # 0.5 + Wed Thu Sat = 3.5 days x 5

    def test_future_days_are_not_costed(self):
        row = self.only(self.cost(first=SUN, last=date(2999, 1, 1))).rows[0]
        self.assertEqual(row[4], D("6"))

    def test_a_monthly_salary_is_spread_over_the_days_allocated_less_absences(self):
        monthly = self.contracted("Salaried")
        Employee.objects.filter(pk=monthly.pk).update(employment_type="permanent")
        profile = monthly.labor_profile
        profile.engagement = "direct"
        profile.contractor = None
        profile.save()
        services.change_rate(profile, effective_from=date(2026, 2, 1), wage_basis="monthly", rate=D("300"),
                             standard_hours=D("8"), overtime_eligible=True)
        self.present()
        Attendance.objects.filter(employee=monthly, date=MON).update(status="absent")
        rows = {r[0]: r for r in self.only(self.cost("worker")).rows}
        salaried = next(r for r in self.only(self.cost("worker")).rows if r[1] == "Salaried")
        self.assertEqual(salaried[6], D("60.000"))                           # 7 calendar days at 300/30 = 10, less 1 absent day

    def test_a_day_worked_on_a_rest_day_is_paid_as_overtime_when_the_rules_say_so(self):
        sitesheet.save_sheet(self.root, self.p1, self.site1, FRI, FRI, {self.key(self.emp, FRI): "present"}, "")
        row = self.only(self.cost()).rows[0]
        self.assertEqual((row[4], row[6]), (D("7"), D("30.000")))             # the Friday counts as a day worked, wages unchanged
        self.set_rules(date(2026, 6, 1), all_hours_on_days_off=False)         # later rules do not change March
        self.assertEqual(self.only(self.cost()).rows[0][6], D("30.000"))

    def test_with_no_overtime_company_or_rules_a_rest_day_is_an_ordinary_paid_day(self):
        sitesheet.save_sheet(self.root, self.p1, self.site1, FRI, FRI, {self.key(self.emp, FRI): "present"}, "")
        from .overtime import OvertimePolicy
        OvertimePolicy.objects.all().delete()
        self.assertEqual(self.only(self.cost()).rows[0][6], D("35.000"))      # no rules: 7 days x 5
        self.set_rules(date(2026, 1, 1), overtime_applies=False)
        self.assertEqual(self.only(self.cost()).rows[0][6], D("35.000"))

    def test_every_grouping_adds_up_to_the_same_wages_and_overtime(self):
        self.approve_ravi_overtime()
        self.contracted("Imran")
        self.present()
        totals = set()
        for group in ("site", "work_order", "trade", "contractor", "worker"):
            t = self.only(self.cost(group))
            totals.add(tuple(t.total[-7:]))
        self.assertEqual(len(totals), 1)

    def test_work_orders_split_the_cost(self):
        wo = WorkOrder.objects.create(project=self.p1, code="WO-1", name="Foundations")
        deployment.release(self.emp, last_day=date(2026, 3, 3))
        deployment.allocate(self.emp, project=self.p1, location=self.site1, work_order=wo, effective_from=WED)
        t = self.only(self.cost("work_order"))
        by = {r[1]: r for r in t.rows}
        self.assertEqual(by["(no work order)"][6], D("15.000"))                # Sun Mon Tue
        self.assertEqual(by["WO-1"][6], D("15.000"))                           # Wed Thu Sat

    def test_grouping_by_trade_contractor_and_worker(self):
        self.contracted("Imran")
        self.present()
        self.assertEqual([r[0] for r in self.only(self.cost("trade")).rows], ["Mason"])
        names = [r[0] for r in self.only(self.cost("contractor")).rows]
        self.assertEqual(names, ["Gulf Manpower", "Direct employees"])
        self.assertEqual(len(self.only(self.cost("worker")).rows), 2)
        self.assertEqual(self.only(self.cost("nonsense")).title, "By site")      # an unknown grouping falls back to site

    def test_filters(self):
        self.contracted("Imran")
        self.present()
        self.assertEqual(len(self.only(self.cost("worker", engagement="contracted")).rows), 1)
        self.assertEqual(len(self.only(self.cost("worker", contractor="none")).rows), 1)       # Ravi has no contractor
        self.assertEqual(len(self.only(self.cost("worker", contractor=str(self.contractor.pk))).rows), 1)
        self.assertEqual(self.only(self.cost("worker", company=str(self.other.pk))).rows, [])

    def test_rejected_overtime_is_ignored_and_pending_is_kept_apart(self):
        self.confirmed({self.hkey(self.emp, MON): "10", self.hkey(self.emp, TUE): "11"})
        self.prepare()
        ot.decide(self.root, OvertimeClaim.objects.filter(date=MON), False, "no")
        row = self.only(self.cost()).rows[0]
        self.assertEqual((row[8], row[9], row[10]), (D("0.000"), D("30.000"), D("2.344")))

    def test_other_companies_see_nothing_and_the_basis_is_always_stated(self):
        outsider = get_user_model().objects.create_user("out")
        rep = reporting.cost_report(outsider, SUN, SAT)
        self.assertEqual((rep.tables[0].rows, rep.tables[0].total), ([], None))
        self.assertTrue(any("estimate" in n for n in rep.notes))
        self.assertTrue(any("weekly off or holiday" in n for n in rep.notes))


class StatementTests(ReportCase):
    def test_amount_payable_is_days_times_rate_plus_approved_overtime(self):
        emp = self.contracted("Imran", rate="6")
        self.present()
        sitesheet.save_sheet(self.root, self.p1, self.site1, SUN, SAT, {}, "present")
        timekeeping.save_hours(self.root, self.p1, self.site1, SUN, SAT, {self.hkey(emp, MON): "10"}, "standard")
        timekeeping.confirm_period(self.root, self.p1, self.site1, SUN, SAT)
        self.prepare()
        ot.decide(self.root, OvertimeClaim.objects.filter(employee=emp), True)
        rep = reporting.contractor_statement(self.user, SUN, SAT, str(self.contractor.pk))
        row = rep.tables[0].rows[0]
        self.assertEqual(row[:5], [emp.employee_no, "Imran", "Mason", D("6"), "6.000"])
        self.assertEqual((row[5], row[7]), (D("36.000"), row[5] + row[6]))
        self.assertEqual(row[6], D("1.875"))                                    # 2 x (6/8) x 1.25
        self.assertEqual(rep.tables[0].total[7], row[7])
        self.assertEqual(rep.tables[1].total[-1], row[7])                       # the by-site table adds up to the same amount

    def test_only_that_contractors_workers_and_only_contracted_ones(self):
        self.contracted("Imran")
        self.contracted("Direct client", contractor=None)
        self.present()
        rep = reporting.contractor_statement(self.user, SUN, SAT, str(self.contractor.pk))
        self.assertEqual([r[1] for r in rep.tables[0].rows], ["Imran"])
        none = reporting.contractor_statement(self.user, SUN, SAT, "none")
        self.assertEqual([r[1] for r in none.tables[0].rows], ["Direct client"])
        self.assertIn("engaged directly", none.subtitle)

    def test_a_changed_rate_shows_as_varies_and_pending_overtime_is_noted_not_paid(self):
        emp = self.contracted("Imran", rate="6")
        services.change_rate(emp.labor_profile, effective_from=WED, wage_basis="daily", rate=D("7"), standard_hours=D("8"),
                             overtime_eligible=True)
        self.present()
        timekeeping.save_hours(self.root, self.p1, self.site1, SUN, SAT, {self.hkey(emp, MON): "10"}, "standard")
        timekeeping.confirm_period(self.root, self.p1, self.site1, SUN, SAT)
        self.prepare()
        rep = reporting.contractor_statement(self.user, SUN, SAT, str(self.contractor.pk))
        row = rep.tables[0].rows[0]
        self.assertEqual(row[4], "varies")
        self.assertEqual(row[5], D("3") * D("6") + D("3") * D("7"))             # Sun Mon Tue at 6, Wed Thu Sat at 7
        self.assertEqual(row[6], D("0.000"))
        self.assertTrue(any("pending approval" in n for n in rep.notes))


@unittest.skipIf(openpyxl is None, "openpyxl is only used here to read the file back")
class ExcelTests(ReportCase):
    def test_the_file_holds_exactly_what_the_screen_shows(self):
        self.approve_ravi_overtime()
        report = self.cost()
        wb = openpyxl.load_workbook(io.BytesIO(reporting.to_xlsx(report)))
        ws = wb["By site"]
        self.assertEqual(ws["A1"].value, "Labor cost")
        self.assertEqual([c.value for c in ws[4]][:4], ["Project", "Project name", "Site", "Workers"])
        table = report.tables[0]
        self.assertEqual([c.value for c in ws[5]][:3], ["P1", "Tower A", "Site One"])
        self.assertEqual(ws["G5"].value, float(table.rows[0][6]))                # wages
        self.assertEqual(ws["G5"].number_format, "#,##0.000")
        self.assertEqual(ws["D6"].value, table.total[3])
        self.assertTrue(ws["A6"].font.b)
        notes = [c.value for row in ws.iter_rows(min_col=1, max_col=1) for c in row if c.value and "estimate" in str(c.value)]
        self.assertTrue(notes)
        self.assertEqual(ws.freeze_panes, "A5")

    def test_one_sheet_per_table(self):
        names = openpyxl.load_workbook(io.BytesIO(reporting.to_xlsx(reporting.deployment_report(self.user, date(2026, 3, 3))))).sheetnames
        self.assertEqual(names, ["By site", "By trade", "By contractor"])
