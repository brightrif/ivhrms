import io
import unittest
from datetime import date

from django.contrib.auth.models import Permission

from apps.labor import deployment, sitesheet, timekeeping
from apps.labor import overtime_services as ot
from apps.labor.overtime import OvertimeClaim

from .test_labor_overtime import OvertimePageCase

try:
    import openpyxl
except ImportError:
    openpyxl = None

SUN, MON, SAT = date(2026, 3, 1), date(2026, 3, 2), date(2026, 3, 7)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ReportPageCase(OvertimePageCase):
    """Ravi: 6 days at 5.000 at P1 / Site One in the week of 1 March, hours 10 and 11, overtime prepared and one day approved."""

    def setUp(self):
        super().setUp()
        self.rules()
        self.post("prepare")
        ot.decide(self.manager, OvertimeClaim.objects.filter(date=MON), True)
        self.cost_q = {"from": "2026-03-01", "to": "2026-03-07", "group": "site"}

    def cost(self, **q):
        return self.client.get(self.url("labor_report_cost"), {**self.cost_q, **q})

    def only_see_headcount(self):
        user = type(self.hr).objects.create_user("headcount")
        user.user_permissions.add(*Permission.objects.filter(content_type__app_label="labor",
                                                             codename__in=["view_laborallocation", "view_laborprofile"]))
        self.hr.company_access.all().update()
        from apps.organization.models import CompanyAccess
        CompanyAccess.objects.create(user=user, company=self.co)
        return user


class AccessTests(ReportPageCase):
    PAGES = ("labor_report_deployment", "labor_report_manpower", "labor_report_cost", "labor_report_contractor")

    def test_needs_login_and_a_right(self):
        self.client.logout()
        for name in self.PAGES:
            self.assertEqual(self.client.get(self.url(name)).status_code, 302, name)
        self.client.force_login(self.nobody)
        for name in self.PAGES:
            self.assertEqual(self.client.get(self.url(name)).status_code, 403, name)

    def test_hr_finance_and_management_can_open_every_report(self):
        for user in (self.hr, self.finance, self.boss):
            self.client.force_login(user)
            for name in self.PAGES:
                self.assertEqual(self.client.get(self.url(name)).status_code, 200, (user.username, name))

    def test_someone_who_sees_headcount_but_not_pay_never_sees_money(self):
        self.client.force_login(self.only_see_headcount())
        for name in ("labor_report_deployment", "labor_report_manpower"):
            self.assertEqual(self.client.get(self.url(name)).status_code, 200, name)
        for name in ("labor_report_cost", "labor_report_contractor"):
            self.assertEqual(self.client.get(self.url(name)).status_code, 403, name)
            self.assertEqual(self.client.get(self.url(name), {"format": "xlsx"}).status_code, 403, name)
        page = self.client.get(self.url("labor_report_deployment"))
        self.assertNotContains(page, "BHD")
        self.assertNotContains(page, ">Cost<")                       # the pill for a report they cannot open is not shown
        self.assertContains(page, ">Manpower by day<")

    def test_the_reports_tab_goes_to_the_first_report_the_person_may_open(self):
        self.client.force_login(self.only_see_headcount())
        self.assertContains(self.client.get(self.url("labor_list")), self.url("labor_report_deployment"))
        pay_only = type(self.hr).objects.create_user("payonly")
        pay_only.user_permissions.add(*Permission.objects.filter(content_type__app_label="labor", codename__in=["view_laborrate", "view_laborprofile"]))
        from apps.organization.models import CompanyAccess
        CompanyAccess.objects.create(user=pay_only, company=self.co)
        self.client.force_login(pay_only)
        self.assertContains(self.client.get(self.url("labor_list")), self.url("labor_report_cost"))

    def test_other_companies_see_no_workers_in_any_report(self):
        self.client.force_login(self.outsider)
        self.assertNotContains(self.cost(), "Ravi")
        self.assertNotContains(self.cost(group="worker"), "Ravi")
        self.assertContains(self.cost(), "Nothing to show")
        dep = self.client.get(self.url("labor_report_deployment"), {"date": "2026-03-03"})
        self.assertNotContains(dep, "Tower A")


class ScreenTests(ReportPageCase):
    def test_deployment(self):
        page = self.client.get(self.url("labor_report_deployment"), {"date": "2026-03-03"})
        self.assertContains(page, "Labor deployment")
        self.assertContains(page, "As of 03 Mar 2026")
        self.assertContains(page, "Tower A")
        self.assertContains(page, "By trade")
        self.assertContains(page, "By contractor")
        self.assertContains(page, "Mason")

    def test_manpower_follows_the_period_controls(self):
        page = self.client.get(self.url("labor_report_manpower"), {"period": "week", "date": "2026-03-04"})
        self.assertEqual(len(page.context["tables"][0]["columns"]), 1 + 7 + 1)
        self.assertContains(page, "P1 / Site One")
        month = self.client.get(self.url("labor_report_manpower"), {"period": "month", "date": "2026-03-04"})
        self.assertEqual(len(month.context["tables"][0]["columns"]), 1 + 31 + 1)
        self.assertIn("date=2026-02-01", month.context["prev_query"])
        default = self.client.get(self.url("labor_report_manpower"))
        self.assertEqual(default.context["period"], "month")

    def test_cost_shows_the_numbers_worked_out_by_hand(self):
        page = self.cost()
        self.assertContains(page, "Labor cost")
        self.assertContains(page, "30.000")                           # 6 days x 5.000
        self.assertContains(page, "1.563")                            # Monday's approved overtime
        self.assertContains(page, "31.563")                           # wages plus approved overtime
        self.assertContains(page, "2.344")                            # Tuesday's pending overtime, kept apart
        self.assertContains(page, "An estimate for cost control")

    def test_every_grouping_and_filter_works_and_unknown_ones_fall_back(self):
        for group in ("site", "work_order", "trade", "contractor", "worker"):
            self.assertEqual(self.cost(group=group).status_code, 200, group)
        self.assertContains(self.cost(group="worker"), "Ravi")
        self.assertContains(self.cost(group="nonsense"), "By site")
        self.assertNotContains(self.cost(engagement="contracted"), "30.000")
        self.assertContains(self.cost(engagement="direct"), "30.000")
        self.assertContains(self.cost(company="999"), "30.000")        # a company the person cannot choose is ignored

    def test_bad_dates_are_handled_politely(self):
        self.assertContains(self.cost(**{"from": "2026-03-07", "to": "2026-03-01"}), "end date is before the start date")
        long = self.cost(**{"from": "2020-01-01", "to": "2026-03-07"})
        self.assertEqual(long.status_code, 200)
        self.assertContains(long, "at most 366 days")
        junk = self.cost(**{"from": "not-a-date", "to": "also-not"})
        self.assertEqual(junk.status_code, 200)

    def test_the_contractor_statement_and_its_choices(self):
        page = self.client.get(self.url("labor_report_contractor"), {"from": "2026-03-01", "to": "2026-03-07"})
        self.assertEqual(page.context["contractor"], str(self.contractor.pk))            # the first contractor by default
        self.assertContains(page, "Contractor statement")
        self.assertContains(page, "Gulf Manpower")
        none = self.client.get(self.url("labor_report_contractor"), {"contractor": "none", "from": "2026-03-01", "to": "2026-03-07"})
        self.assertEqual(none.context["contractor"], "none")
        bogus = self.client.get(self.url("labor_report_contractor"), {"contractor": "99999"})
        self.assertEqual(bogus.context["contractor"], str(self.contractor.pk))

    def test_each_report_offers_the_excel_download_with_the_same_choices(self):
        page = self.cost(group="trade")
        self.assertContains(page, "Export to Excel")
        self.assertIn("group=trade", page.context["export_query"])
        self.assertIn("format=xlsx", page.context["export_query"])


@unittest.skipIf(openpyxl is None, "openpyxl is only used here to read the downloaded file back")
class ExcelDownloadTests(ReportPageCase):
    def test_the_download_is_a_real_workbook_with_the_same_numbers(self):
        r = self.cost(format="xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], XLSX)
        self.assertEqual(r["Content-Disposition"], 'attachment; filename="labor-cost-2026-03-01_2026-03-07-by-site.xlsx"')
        ws = openpyxl.load_workbook(io.BytesIO(r.content))["By site"]
        self.assertEqual(ws["A1"].value, "Labor cost")
        values = [c.value for row in ws.iter_rows() for c in row]
        self.assertIn("Tower A", values)
        self.assertIn(30.0, values)
        self.assertIn(31.563, values)
        screen = self.cost()
        self.assertEqual(screen.context["tables"][0]["total"][6][0], "30.000")

    def test_every_report_can_be_downloaded(self):
        for name, extra in (("labor_report_deployment", {"date": "2026-03-03"}), ("labor_report_manpower", {"period": "week", "date": "2026-03-04"}),
                            ("labor_report_cost", {"from": "2026-03-01", "to": "2026-03-07"}),
                            ("labor_report_contractor", {"from": "2026-03-01", "to": "2026-03-07"})):
            r = self.client.get(self.url(name), {**extra, "format": "xlsx"})
            self.assertEqual(r.status_code, 200, name)
            self.assertEqual(r["Content-Type"], XLSX, name)
            self.assertTrue(openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames, name)

    def test_the_download_respects_who_may_see_what(self):
        self.client.force_login(self.outsider)
        ws = openpyxl.load_workbook(io.BytesIO(self.cost(format="xlsx").content))["By site"]
        self.assertNotIn("Tower A", [c.value for row in ws.iter_rows() for c in row])

    def test_the_templates_use_only_the_current_styles(self):
        import re
        from pathlib import Path
        legacy = re.compile(r'class="(pill|muted|err|ok|big|head|cards|chips|scroll|pager|filters|inline-form|form|card narrow)[ "]')
        text = (Path(__file__).resolve().parent.parent / "templates" / "web" / "labor" / "report.html").read_text(encoding="utf-8")
        self.assertIsNone(legacy.search(text))
        self.assertNotIn("<table>", text)
